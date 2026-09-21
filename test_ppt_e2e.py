import time
import os
import sys

try:
    import win32com.client
except ImportError:
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "pywin32", "-q"])
    import win32com.client

SAVE_PATH = r"E:\agent_test\ppt_result.pptx"

def main():
    total_start = time.time()
    
    print("=" * 60, flush=True)
    print("WPS PPT E2E Test (COM)", flush=True)
    print("=" * 60, flush=True)
    
    wpp = None
    
    try:
        # Step 1: 启动WPS演示
        t = time.time()
        wpp = win32com.client.Dispatch("KWpp.Application")
        wpp.Visible = True
        time.sleep(2)
        print(f"[Step 1] 启动WPS演示             {time.time()-t:.2f}s", flush=True)
        
        # Step 2: 新建演示文稿
        t = time.time()
        pres = wpp.Presentations.Add()
        time.sleep(1)
        print(f"[Step 2] 新建演示文稿            {time.time()-t:.2f}s", flush=True)
        
        # Step 3: 添加标题幻灯片
        t = time.time()
        slide = pres.Slides.Add(1, 1)  # 1 = ppLayoutTitle
        time.sleep(0.5)
        print(f"[Step 3] 添加标题幻灯片          {time.time()-t:.2f}s", flush=True)
        
        # Step 4: 输入标题
        t = time.time()
        slide.Shapes(1).TextFrame.TextRange.Text = "AI Agent Presentation"
        print(f"[Step 4] 输入标题               {time.time()-t:.2f}s", flush=True)
        
        # Step 5: 保存
        t = time.time()
        os.makedirs(os.path.dirname(SAVE_PATH), exist_ok=True)
        pres.SaveAs(SAVE_PATH)
        print(f"[Step 5] 保存文件               {time.time()-t:.2f}s", flush=True)
        
        # Step 6: 关闭
        t = time.time()
        try:
            pres.Close()
        except:
            pass
        try:
            wpp.Quit()
        except:
            pass
        print(f"[Step 6] 关闭WPS演示            {time.time()-t:.2f}s", flush=True)
        
        # Step 7: 验证文件
        t = time.time()
        exists = os.path.exists(SAVE_PATH)
        size = os.path.getsize(SAVE_PATH) if exists else 0
        print(f"[Step 7] 文件验证: {exists}, {size}B   {time.time()-t:.2f}s", flush=True)
        
        # Step 8: 读取内容验证
        t = time.time()
        has_title = False
        if exists:
            from pptx import Presentation
            prs = Presentation(SAVE_PATH)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        if "AI Agent Presentation" in shape.text:
                            has_title = True
                            break
                if has_title:
                    break
        print(f"[Step 8] 内容验证: {'包含标题' if has_title else '标题缺失'}   {time.time()-t:.2f}s", flush=True)
        
        total_time = time.time() - total_start
        
        print("\n" + "=" * 60, flush=True)
        print(f"总耗时: {total_time:.2f}s", flush=True)
        print(f"文件: {SAVE_PATH}", flush=True)
        print(f"存在: {exists}, 大小: {size}B", flush=True)
        print(f"标题验证: {has_title}", flush=True)
        
        passed = exists and size > 0 and has_title and total_time < 30
        print(f"\n{'✓ TEST PASSED' if passed else '✗ TEST FAILED'}", flush=True)
        return passed
        
    except Exception as e:
        print(f"\n✗ ERROR: {e}", flush=True)
        try:
            if wpp: wpp.Quit()
        except: pass
        return False

if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
